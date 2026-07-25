"""Rendu XLSX du questionnaire d'un AO, via `openpyxl`.

Contrairement au PDF/DOCX (pensés comme un rapport imprimable), le XLSX est un
tableau simple et exploitable (filtrable, triable) : une ligne par question.
"""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from mauricette.api.export.theme import MARINE
from mauricette.application.cas_usage.exporter_reponses_appel_offre import ExportReponses

_ENTETES = ["Référentiel", "Section", "Question", "Obligatoire", "Réponse", "Sources"]
_LARGEURS = [22, 22, 40, 12, 50, 30]


def generer_xlsx(export: ExportReponses) -> bytes:
    """Génère le classeur XLSX du questionnaire complet (une ligne par question)."""
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Réponses"

    couleur_entete = f"FF{MARINE.upper()}"
    fond_entete = PatternFill(start_color=couleur_entete, end_color=couleur_entete, fill_type="solid")
    police_entete = Font(color="FFFFFF", bold=True)
    for colonne, libelle in enumerate(_ENTETES, start=1):
        cellule = feuille.cell(row=1, column=colonne, value=libelle)
        cellule.fill = fond_entete
        cellule.font = police_entete

    for colonne, largeur in enumerate(_LARGEURS, start=1):
        feuille.column_dimensions[get_column_letter(colonne)].width = largeur

    ligne = 2
    alignement_haut = Alignment(vertical="top", wrap_text=True)
    for referentiel in export.referentiels:
        for section in referentiel.sections:
            for question in section.questions:
                valeurs = [
                    referentiel.nom,
                    section.nom,
                    question.question,
                    "Oui" if question.obligatoire else "Non",
                    question.contenu or "Non renseigné",
                    " · ".join(question.citations),
                ]
                for colonne, valeur in enumerate(valeurs, start=1):
                    cellule = feuille.cell(row=ligne, column=colonne, value=valeur)
                    cellule.alignment = alignement_haut
                ligne += 1

    tampon = io.BytesIO()
    classeur.save(tampon)
    return tampon.getvalue()
