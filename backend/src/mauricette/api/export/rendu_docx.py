"""Rendu DOCX du questionnaire d'un AO, via `python-docx`.

Reprend volontairement la même hiérarchie visuelle que `rendu_pdf.py`
(bandeau, H1 référentiel, H2 section + ligne colorée, question en gras,
réponse ou "Non renseigné", citations en petit texte gris) pour que les
deux formats se ressemblent, comme demandé.
"""

from __future__ import annotations

import io

from docx import Document as DocumentDocx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from mauricette.api.export.theme import CYAN, GRIS_DISCRET, MARINE
from mauricette.application.cas_usage.exporter_reponses_appel_offre import ExportReponses


def _couleur(hexa: str) -> RGBColor:
    return RGBColor.from_string(hexa)


def _ajouter_bordure_basse(paragraphe, couleur_hex: str, taille: int = 12) -> None:
    """Ajoute une bordure basse à un paragraphe : python-docx n'a pas d'API de haut
    niveau pour ça, il faut manipuler l'OXML directement (`w:pBdr`/`w:bottom`)."""
    pPr = paragraphe._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(taille))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), couleur_hex)
    pBdr.append(bottom)
    pPr.append(pBdr)


def generer_docx(export: ExportReponses) -> bytes:
    """Génère le DOCX du questionnaire complet (une question par bloc, "Non renseigné" si vide)."""
    document = DocumentDocx()

    # Bandeau marine : un tableau 1x1 sur toute la largeur, cellule ombrée.
    tableau_bandeau = document.add_table(rows=1, cols=1)
    cellule = tableau_bandeau.rows[0].cells[0]
    ombrage = OxmlElement("w:shd")
    ombrage.set(qn("w:fill"), MARINE)
    cellule._tc.get_or_add_tcPr().append(ombrage)
    paragraphe_titre = cellule.paragraphs[0]
    run_titre = paragraphe_titre.add_run(export.nom_appel_offre)
    run_titre.font.color.rgb = _couleur("FFFFFF")
    run_titre.font.bold = True
    run_titre.font.size = Pt(18)
    document.add_paragraph()

    for referentiel in export.referentiels:
        titre_h1 = document.add_heading(referentiel.nom, level=1)
        for run in titre_h1.runs:
            run.font.color.rgb = _couleur(MARINE)

        for section in referentiel.sections:
            titre_h2 = document.add_heading(section.nom, level=2)
            for run in titre_h2.runs:
                run.font.color.rgb = _couleur(MARINE)
            _ajouter_bordure_basse(titre_h2, CYAN)

            for question in section.questions:
                paragraphe_question = document.add_paragraph()
                run_question = paragraphe_question.add_run(question.question)
                run_question.font.bold = True

                if question.obligatoire:
                    paragraphe_obligatoire = document.add_paragraph()
                    run_obligatoire = paragraphe_obligatoire.add_run("Obligatoire")
                    run_obligatoire.font.italic = True
                    run_obligatoire.font.size = Pt(8)
                    run_obligatoire.font.color.rgb = _couleur(CYAN)

                paragraphe_reponse = document.add_paragraph()
                run_reponse = paragraphe_reponse.add_run(question.contenu or "Non renseigné")
                if not question.contenu:
                    run_reponse.font.italic = True
                    run_reponse.font.color.rgb = _couleur(GRIS_DISCRET)

                if question.citations:
                    paragraphe_citations = document.add_paragraph()
                    run_citations = paragraphe_citations.add_run(" · ".join(question.citations))
                    run_citations.font.size = Pt(8)
                    run_citations.font.color.rgb = _couleur(GRIS_DISCRET)

    tampon = io.BytesIO()
    document.save(tampon)
    return tampon.getvalue()
